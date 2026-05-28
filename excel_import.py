"""
excel_import.py
Reads an existing schedule Excel file and returns employee states for continuity.
Handles: merged cells, datetime objects for month header, datetime.time for shift cells.
"""

import re
import calendar
from datetime import date, timedelta, datetime, time as dt_time

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

MONTH_NAMES = {
    "enero":1,"febrero":2,"marzo":3,"abril":4,
    "mayo":5,"junio":6,"julio":7,"agosto":8,
    "septiembre":9,"octubre":10,"noviembre":11,"diciembre":12
}

# ── Value normaliser ──────────────────────────────────────────────────────────

def _normalise(raw) -> str:
    """Convert a cell value to one of: AM | PM | NIGHT | V | D | L
    V = vacaciones (día libre, se registra aparte)
    D = devolución (día trabajado normal, se trata como el turno vigente)
    """
    if raw is None:
        return "L"

    # datetime.time objects (most common in this Excel)
    if isinstance(raw, dt_time):
        if raw.hour == 14:  return "AM"
        if raw.hour == 6:   return "PM"
        if raw.hour in (19, 22): return "NIGHT"  # 19:00 y 22:00 → NIGHT
        if raw.hour == 8:   return "PM"   # 8:00/8:30 → PM (turno mañana)
        return "L"

    # datetime objects (the header cell)
    if isinstance(raw, datetime):
        return "L"

    s = str(raw).strip().lower()
    if s in ("", "l", "libre", "-"):
        return "L"
    if s == "v":   return "V"   # vacaciones
    if s == "d":   return "D"   # devolución
    if s == "ad":  return "D"   # AD = ajuste/devolución también
    s2 = s.replace(" ", "").replace(":", "")
    if s2 in ("1400", "14"):         return "AM"
    if s2 in ("600", "0600", "830", "0830", "800", "0800"): return "PM"
    if s2 in ("2200", "22", "1900", "19"):  return "NIGHT"
    if s2 in ("1530", "15:30"):      return "AM"  # T8 15:30 → AM aproximado
    return "L"


def _shift_from_fill(cell) -> str | None:
    try:
        rgb = cell.fill.fgColor.rgb.upper()  # AARRGGBB
        r, g, b = int(rgb[2:4],16), int(rgb[4:6],16), int(rgb[6:8],16)
        if b > 150 and g > 150 and r < 100: return "NIGHT"
        if b > 150 and r < 100 and g < 100: return "AM"
        if r > 150 and g > 80  and b < 80:  return "PM"
    except Exception:
        pass
    return None


# ── Main import function ──────────────────────────────────────────────────────

def import_excel(filepath: str) -> dict:
    if not HAS_OPENPYXL:
        raise RuntimeError("openpyxl is not installed. Run: pip install openpyxl")

    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active

    warnings  = []
    year = month = None
    day_cols  = {}   # col_index → day_number
    name_col  = None
    employees = []
    in_specialists = False

    rows = list(ws.iter_rows(values_only=False))

    # ── 1. Find month/year ────────────────────────────────────────────────────
    # Priority: merged cells first (they often hold the title)
    candidates = []
    for mr in ws.merged_cells.ranges:
        tl = ws.cell(mr.min_row, mr.min_col)
        if tl.value is not None:
            candidates.append(tl.value)
    for row in rows:
        for cell in row:
            if cell.value is not None:
                candidates.append(cell.value)

    for val in candidates:
        # Case 1: it's a datetime → month & year from the object directly
        if isinstance(val, datetime):
            month = val.month
            year  = val.year
            break
        t = str(val).strip().lower()
        # Case 2: "abril 2026" (mes primero)
        m = re.search(
            r"(enero|febrero|marzo|abril|mayo|junio|julio|agosto|"
            r"septiembre|octubre|noviembre|diciembre)[^\d]*(\d{4})", t
        )
        if m:
            month = MONTH_NAMES[m.group(1)]
            year  = int(m.group(2))
            break
        # Case 3: "2026 abril" (año primero)
        m2 = re.search(
            r"(\d{4})[^\d]*(enero|febrero|marzo|abril|mayo|junio|julio|agosto|"
            r"septiembre|octubre|noviembre|diciembre)", t
        )
        if m2:
            year  = int(m2.group(1))
            month = MONTH_NAMES[m2.group(2)]
            break

    if not year:
        for val in candidates:
            m = re.search(r"\b(20\d{2})\b", str(val))
            if m:
                year = int(m.group(1))
                break

    # Fallback: extraer del nombre del archivo
    import os
    fname_lower = os.path.basename(filepath).lower()
    if not year:
        ym = re.search(r"\b(20\d{2})\b", fname_lower)
        if ym:
            year = int(ym.group(1))
    if not month:
        mm = re.search(
            r"(enero|febrero|marzo|abril|mayo|junio|julio|agosto|"
            r"septiembre|octubre|noviembre|diciembre)", fname_lower
        )
        if mm:
            month = MONTH_NAMES[mm.group(1)]

    if not year:
        raise ValueError("No se encontró el año en el archivo Excel ni en el nombre del archivo.")
    if not month:
        raise ValueError("No se encontró el mes en el archivo Excel ni en el nombre del archivo.")

    num_days = calendar.monthrange(year, month)[1]

    # ── 2. Find day-numbers row ───────────────────────────────────────────────
    day_row_idx = None
    for ri, row in enumerate(rows):
        # Collect numeric values (int, float, or string digits)
        nums = []
        for cell in row:
            v = cell.value
            if isinstance(v, (int, float)) and 1 <= int(v) <= 31:
                nums.append((cell.column, int(v)))
            elif isinstance(v, str) and v.strip().isdigit():
                n = int(v.strip())
                if 1 <= n <= 31:
                    nums.append((cell.column, n))

        if len(nums) >= num_days - 2:
            day_row_idx = ri
            day_cols    = {col: day for col, day in nums}
            min_day_col = min(day_cols.keys())
            name_col    = min_day_col - 1
            break

    if day_row_idx is None:
        raise ValueError("No se encontró la fila de días en el Excel.")

    # ── 3. Parse employee rows ────────────────────────────────────────────────
    for ri in range(day_row_idx + 1, len(rows)):
        row = rows[ri]

        # Find name cell
        name_cell = next((c for c in row if c.column == name_col), None)
        if name_cell is None:
            continue

        raw_name = str(name_cell.value or "").strip()
        if not raw_name:
            continue

        raw_lower = raw_name.lower()
        if any(kw in raw_lower for kw in ("especialista", "separator", "---")):
            in_specialists = True
            continue

        # Clean name
        clean_name = re.sub(r"\s*\(D:[^)]*\)", "", raw_name).strip()
        if not clean_name or clean_name.lower() in ("v", ""):
            continue

        emp_type = "especialista" if in_specialists else "rotativo"

        # Read assignments — track V (vacation) and D (devolution) separately
        assignments  = {}
        vacation_days_list = []   # day numbers marked as V
        raw_shifts   = {}         # before V/D resolution

        for cell in row:
            if cell.column not in day_cols:
                continue
            day_num = day_cols[cell.column]
            shift   = _normalise(cell.value)
            if shift == "L":
                cs = _shift_from_fill(cell)
                if cs:
                    shift = cs
            raw_shifts[day_num] = shift

        for d in range(1, num_days + 1):
            raw_shifts.setdefault(d, "L")

        # Resolve V and D:
        # V → libre en el horario + registrar como vacación
        # D → usar el turno trabajado ese día (ya viene como AM/PM/NIGHT desde _normalise
        #      si tenía hora; si quedó como "D" puro, tratarlo como el turno más cercano)
        for d in range(1, num_days + 1):
            s = raw_shifts[d]
            if s == "V":
                vacation_days_list.append(d)
                assignments[str(d)] = "L"
            elif s == "D":
                # Devolución: buscar turno vigente del contexto cercano
                # Mirar días anteriores para encontrar el turno en curso
                ctx = "L"
                for dd in range(d - 1, max(0, d - 5), -1):
                    prev = raw_shifts.get(dd, "L")
                    if prev not in ("L", "V", "D"):
                        ctx = prev
                        break
                assignments[str(d)] = ctx if ctx != "L" else "AM"  # fallback AM
            else:
                assignments[str(d)] = s

        # Last worked shift/date
        last_shift = last_date = None
        for d in range(num_days, 0, -1):
            v = assignments.get(str(d), "L")
            if v != "L":
                last_shift = v
                last_date  = date(year, month, d).isoformat()
                break

        # Count trailing rest days (rest_served)
        rest_served = 0
        for d in range(num_days, 0, -1):
            v = assignments.get(str(d), "L")
            if v == "L":
                rest_served += 1
            else:
                break

        # Count block_remaining: how many more work days needed to complete last block
        # Walk backwards: skip trailing L, then count consecutive last_shift days
        block_done = 0
        if last_shift:
            found_block = False
            for d in range(num_days, 0, -1):
                v = assignments.get(str(d), "L")
                if v == last_shift:
                    found_block = True
                    block_done += 1
                elif found_block:
                    break  # hit something different after finding the block
                # else: still in trailing rest, keep going
        block_pref = {"PM": 4, "AM": 3, "NIGHT": 3}
        block_remaining = max(0, block_pref.get(last_shift, 0) - block_done) if last_shift else 0

        # Specialist starting shift
        spec_start = None
        if emp_type == "especialista":
            for d in range(1, num_days + 1):
                v = assignments.get(str(d), "L")
                if v in ("AM", "PM"):
                    spec_start = v
                    break

        employees.append({
            "name":             clean_name,
            "type":             emp_type,
            "specialist_start": spec_start,
            "assignments":      assignments,
            "vacation_days":    vacation_days_list,
            "last_shift":       last_shift,
            "last_work_date":   last_date,
            "rest_served":      rest_served,
            "block_remaining":  block_remaining,
        })

    if not employees:
        raise ValueError("No se encontraron empleados en el archivo.")

    return {"year": year, "month": month, "employees": employees, "warnings": warnings}
